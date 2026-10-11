"""Generates the test fixtures in tests/fixtures and the OCR golden images in
src/ocr/fixtures. Run via scripts/gen_fixtures.sh (inside the toolbox image).

Every document carries known marker strings that the integration tests and
the smoke script assert on.

    generate.py                 # everything
    generate.py odf images      # only these groups (see GROUPS at the end)
"""

import io
import os
import subprocess
import sys
import tempfile
import zlib

from PIL import Image, ImageDraw, ImageFont, features

OUT = "/work/tests/fixtures"
GOLDEN = "/work/src/ocr/fixtures"
NOTO = "/usr/share/fonts/truetype/noto"
FONT_LATIN = f"{NOTO}/NotoSans-Regular.ttf"
FONT_DEVA = f"{NOTO}/NotoSansDevanagari-Regular.ttf"
FONT_ARAB = f"{NOTO}/NotoNaskhArabic-Regular.ttf"
DEJAVU = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"

os.makedirs(OUT, exist_ok=True)
os.makedirs(GOLDEN, exist_ok=True)
assert features.check("raqm"), "Pillow needs libraqm for complex-script shaping"


def out(name):
    return os.path.join(OUT, name)


# --------------------------------------------------------------------------
# Office (OOXML) + legacy formats via LibreOffice
# --------------------------------------------------------------------------

def make_docx(path, title, paragraphs, table=None):
    import docx

    d = docx.Document()
    d.core_properties.title = title
    d.core_properties.author = "Oxide Fixtures"
    d.add_heading(title, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    d.save(path)


def make_office():
    make_docx(
        out("sample.docx"),
        "Quarterly Report",
        ["Oxide fixture marker: docx-alpha.", "Revenue grew in every region this quarter."],
        [["Region", "Revenue"], ["North", "1200"], ["South", "950"]],
    )

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1"
    ws.append(["Region", "Revenue", "Marker"])
    ws.append(["North", 1200, "xlsx-alpha"])
    ws.append(["South", 950, ""])
    ws2 = wb.create_sheet("Q2")
    ws2.append(["Region", "Revenue"])
    ws2.append(["East", 700])
    wb.properties.title = "Revenue Workbook"
    wb.save(out("sample.xlsx"))

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i, (title, body) in enumerate(
        [("Roadmap 2026", "Oxide fixture marker: pptx-alpha"), ("Next Steps", "Ship the OCR image")]
    ):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
    prs.save(out("sample.pptx"))

    # Legacy binary formats: let LibreOffice write Word 97 / Excel 97 / PowerPoint 97.
    with tempfile.TemporaryDirectory() as tmp:
        for src, filt, ext in [
            ("sample.docx", "MS Word 97", "doc"),
            ("sample.xlsx", "MS Excel 97", "xls"),
            ("sample.pptx", "MS PowerPoint 97", "ppt"),
        ]:
            subprocess.run(
                ["soffice", "--headless", f"-env:UserInstallation=file://{tmp}/lo",
                 "--convert-to", f"{ext}:{filt}", "--outdir", OUT, out(src)],
                check=True, capture_output=True,
            )

    # Encrypted OOXML (CFB with EncryptionInfo + EncryptedPackage).
    import msoffcrypto.format.ooxml as ooxml

    with open(out("sample.docx"), "rb") as f_in, open(out("encrypted.docx"), "wb") as f_out:
        ooxml.OOXMLFile(f_in).encrypt("secret-password", f_out)


def make_odf():
    """OpenDocument versions of the OOXML samples, saved by LibreOffice, so
    they carry the same markers (docx-alpha, xlsx-alpha, pptx-alpha)."""
    with tempfile.TemporaryDirectory() as tmp:
        for src, ext in [("sample.docx", "odt"), ("sample.xlsx", "ods"), ("sample.pptx", "odp")]:
            subprocess.run(
                ["soffice", "--headless", f"-env:UserInstallation=file://{tmp}/lo",
                 "--convert-to", ext, "--outdir", OUT, out(src)],
                check=True, capture_output=True,
            )


# --------------------------------------------------------------------------
# PDFs
# --------------------------------------------------------------------------

def make_native_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(out("native.pdf"), pagesize=A4)
    c.setTitle("Native Fixture")
    c.setAuthor("Oxide Fixtures")
    for i, marker in enumerate(["page-one-alpha", "page-two-beta", "page-three-gamma"], 1):
        c.setFont("Helvetica-Bold", 20)
        c.drawString(72, 760, f"Chapter {i}")
        c.setFont("Helvetica", 12)
        y = 720
        for line in [
            f"Oxide fixture marker: {marker}.",
            "This page has a normal text layer and needs no OCR.",
            "The quick brown fox jumps over the lazy dog.",
        ]:
            c.drawString(72, y, line)
            y -= 18
        c.showPage()
    c.save()


def lo_pdf_from_docx(name, title, paragraphs):
    """Complex-script native PDFs: LibreOffice shapes Devanagari/Arabic properly."""
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, name.replace(".pdf", ".docx"))
        make_docx(src, title, paragraphs)
        subprocess.run(
            ["soffice", "--headless", f"-env:UserInstallation=file://{tmp}/lo",
             "--convert-to", "pdf", "--outdir", OUT, src],
            check=True, capture_output=True,
        )


def make_arabic_forms_pdf():
    """Arabic drawn as *presentation forms* (pre-shaped glyph code points), as
    many legacy producers emit it; the API must normalize them to base letters."""
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(TTFont("DejaVu", DEJAVU))
    c = canvas.Canvas(out("arabic_forms.pdf"), pagesize=A4)
    c.setFont("DejaVu", 22)
    for i, line in enumerate(["مرحبا بالعالم", "السلام عليكم"]):
        shaped = get_display(arabic_reshaper.reshape(line))
        assert any(0xFE70 <= ord(ch) <= 0xFEFF for ch in shaped), "expected presentation forms"
        c.drawRightString(520, 760 - i * 40, shaped)
    c.showPage()
    c.save()


def render_lines(lines, font_path, size, width, height, rtl=False, margin=120):
    img = Image.new("L", (width, height), 255)
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(font_path, size, layout_engine=ImageFont.Layout.RAQM)
    y = margin
    for line in lines:
        kw = {"direction": "rtl", "language": "ar"} if rtl else {}
        box = d.textbbox((0, 0), line, font=font, **kw)
        w = box[2] - box[0]
        x = width - margin - w if rtl else margin
        d.text((x, y), line, font=font, fill=0, **kw)
        y += int(size * 1.9)
    return img


def image_pdf(path, img, dpi=200):
    img.save(path, "PDF", resolution=dpi)


A4_200 = (1654, 2339)  # A4 at 200 DPI


def make_scans():
    image_pdf(out("scan_en.pdf"), render_lines(
        ["Scanned Invoice 4471", "Total amount due: 1,250.00", "Payment terms: thirty days",
         "Thank you for your business"], FONT_LATIN, 56, *A4_200))
    image_pdf(out("scan_hi.pdf"), render_lines(
        ["भारत का संविधान", "हम भारत के लोग", "न्याय स्वतंत्रता समता", "बंधुता बढ़ाने के लिए"],
        FONT_DEVA, 60, *A4_200))
    image_pdf(out("scan_ar.pdf"), render_lines(
        ["مرحبا بالعالم", "السلام عليكم ورحمة الله", "هذه وثيقة ممسوحة ضوئيا", "شكرا لكم"],
        FONT_ARAB, 60, *A4_200, rtl=True))


def make_scan_with_overlay():
    """A Hindi scan saved via a browser's print dialog: the page is the
    image plus a Latin text layer (date, title, URL, page number) that has
    nothing to do with the scan's language."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    scan = render_lines(
        ["भारत का संविधान", "हम भारत के लोग भारत को एक", "संपूर्ण प्रभुत्व संपन्न समाजवादी",
         "पंथनिरपेक्ष लोकतंत्रात्मक गणराज्य", "बनाने के लिए तथा उसके समस्त",
         "नागरिकों को सामाजिक आर्थिक और", "राजनैतिक न्याय प्राप्त कराने के लिए"],
        FONT_DEVA, 60, *A4_200)
    buf = io.BytesIO()
    scan.save(buf, "PNG")
    buf.seek(0)
    w, h = A4
    c = canvas.Canvas(out("scan_hi_overlay.pdf"), pagesize=A4)
    c.drawImage(ImageReader(buf), 36, 36, width=w - 72, height=h - 72)
    c.setFont("Helvetica", 8)
    c.drawString(24, h - 18, "08/10/2026, 17:48")
    c.drawCentredString(w / 2, h - 18, "scanned-constitution-page-1.jpg (1654x2339)")
    c.drawString(24, 14, "https://images.example.com/uploads/documents/scanned-constitution-page-1.jpg")
    c.drawRightString(w - 24, 14, "1/1")
    c.showPage()
    c.save()


def make_scans_with_ocr_layer():
    """A scanner's output: a full-page scan under an invisible (Tr 3) text
    layer from the scanner's own OCR. One layer is right; the other is what
    a scanner writes when it OCRs handwriting, and must be OCR'd again."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    lines = ["Scanned Invoice 4471", "Total amount due: 1,250.00", "Payment terms: thirty days",
             "Thank you for your business"]
    junk = ["IxO'C't. oA cxAtl Voraion.", "inoNie, (le.r'A^cc:^ Wt. rre.n^V",
            "r oyAt. oo^ ^ vVs (AeA\\J<xV\\vje.", "(2or\\)<xVur€. cx.aA 'Vo^'S^oHj"]
    buf = io.BytesIO()
    render_lines(lines, FONT_LATIN, 56, *A4_200).save(buf, "PNG")
    w, h = A4
    for name, layer in (("scan_en_ocr_layer.pdf", lines), ("scan_en_junk_layer.pdf", junk)):
        buf.seek(0)
        c = canvas.Canvas(out(name), pagesize=A4)
        c.drawImage(ImageReader(buf), 0, 0, width=w, height=h)
        # render_lines at 200 DPI: 120 px margin, 56 px text, 1.9 line pitch.
        t = c.beginText(120 * 0.36, h - 120 * 0.36 - 56 * 0.36)
        t.setFont("Helvetica", 56 * 0.36)
        t.setTextRenderMode(3)
        t.setLeading(56 * 1.9 * 0.36)
        for line in layer:
            t.textLine(line)
        c.drawText(t)
        c.showPage()
        c.save()


def make_hybrid():
    """Native text plus an embedded image carrying extra text (a stamp)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    stamp = render_lines(["APPROVED BY FINANCE", "Reference 99812"], FONT_LATIN, 64, 1400, 520, margin=80)
    buf = io.BytesIO()
    stamp.save(buf, "PNG")
    buf.seek(0)
    c = canvas.Canvas(out("hybrid.pdf"), pagesize=A4)
    c.setFont("Helvetica", 12)
    y = 780
    for line in ["Purchase order 5521", "Oxide fixture marker: hybrid-native.",
                 "The stamp below is an image, not text."]:
        c.drawString(72, y, line)
        y -= 18
    c.drawImage(ImageReader(buf), 72, 380, width=450, height=167)
    c.showPage()
    c.save()


def make_legacy_font_pdf():
    """A Kruti Dev-style page: the font is *named* like the legacy Hindi
    family and the text is Kruti Dev's Latin-codepoint encoding of Hindi.
    (Real Kruti Dev fonts aren't redistributable; this exercises routing.)"""
    from fontTools.ttLib import TTFont as FTFont
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    with tempfile.TemporaryDirectory() as tmp:
        f = FTFont(DEJAVU)
        for rec in f["name"].names:
            if rec.nameID in (1, 3, 4, 6, 16):
                rec.string = "KrutiDev010" if rec.nameID == 6 else "Kruti Dev 010"
        renamed = os.path.join(tmp, "KrutiDev010.ttf")
        f.save(renamed)
        pdfmetrics.registerFont(TTFont("KrutiDev010", renamed))
        c = canvas.Canvas(out("legacy_krutidev.pdf"), pagesize=A4)
        c.setFont("KrutiDev010", 20)
        c.drawString(72, 760, "Hkkjr dk lafo/kku")
        c.drawString(72, 720, "ge Hkkjr ds yksx")
        c.showPage()
        c.save()


def make_encrypted_pdfs():
    src = out("native.pdf")
    # Owner password only (empty user password): opens without a password.
    subprocess.run(["qpdf", "--encrypt", "", "owner-pass", "256", "--", src, out("owner_only.pdf")], check=True)
    subprocess.run(["qpdf", "--allow-weak-crypto", "--encrypt", "", "owner-pass", "128", "--use-aes=n", "--",
                    src, out("owner_only_rc4.pdf")], check=True)
    # User password required: truly locked.
    subprocess.run(["qpdf", "--encrypt", "user-pass", "owner-pass", "256", "--", src, out("locked.pdf")], check=True)


def make_nested_pdf(depth=50_000):
    """A syntactically plausible PDF whose page dictionary holds an absurdly
    deep array: parsers must not overflow their stack."""
    nested = b"[" * depth + b"]" * depth
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Junk " + nested + b" >>",
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(body))
        body += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        body += f"{off:010d} 00000 n \n".encode()
    body += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    with open(out("nested.pdf"), "wb") as f:
        f.write(body)


def write_objstm_pdf(name, packed, loose, rows, trailer, update=()):
    """PDF 1.5+ writer for the misfiled-xref fixtures. `packed` become
    objects 1..n inside an ObjStm, `loose` are (key, header, stream payload or
    None) written as they are, `update` likewise but after the ObjStm, the way
    an incremental update appends them; the ObjStm and the xref stream take
    the next two numbers. `rows(at)` gets every object's byte offset and
    returns the xref stream's (type, field 2, field 3) rows, broken as the
    fixture needs."""
    offs, data = [], b""
    for o in packed:
        offs.append(len(data))
        data += o + b"\n"
    head = " ".join(f"{n} {off}" for n, off in enumerate(offs, 1)).encode() + b"\n"
    objstm = zlib.compress(head + data)

    body = b"%PDF-1.7\n%\x81\x81\x81\x81\n"
    at = {}

    def obj(key, header, payload=None):
        nonlocal body
        at[key] = len(body)
        body += header
        if payload is not None:
            body += b"\nstream\n" + payload + b"\nendstream"
        body += b"\nendobj\n"

    for key, header, payload in loose:
        obj(key, header, payload)
    stm = max(k for k, _, _ in [*loose, *update] if isinstance(k, int)) + 1
    obj(stm, f"{stm} 0 obj\n<< /Type /ObjStm /N {len(packed)} /First {len(head)} /Filter /FlateDecode"
             f" /Length {len(objstm)} >>".encode(), objstm)
    for key, header, payload in update:
        obj(key, header, payload)
    at[stm + 1] = len(body)
    table = rows(at)
    xref = zlib.compress(b"".join(bytes([t]) + a.to_bytes(2, "big") + b.to_bytes(2, "big") for t, a, b in table))
    obj(stm + 1, f"{stm + 1} 0 obj\n<< /Type /XRef /Size {stm + 2} {trailer} /W [1 2 2] /Index [0 {len(table)}]"
                 f" /Filter /FlateDecode /Length {len(xref)} >>".encode(), xref)
    body += f"startxref\n{at[stm + 1]}\n%%EOF\n".encode()
    with open(out(name), "wb") as f:
        f.write(body)


def page_content(text):
    content = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET\nq 20 0 0 20 72 640 cm /Im1 Do Q\n".encode()
    return f"<< /Length {len(content)} >>".encode(), content


def make_misfiled_xref_pdfs():
    # One stray xref-stream row: the rewritten Catalog (`1 1 obj`) got a
    # second entry instead of replacing the first, so every later row
    # describes the object before it and the ObjStm's row (8) lands on an
    # image XObject. Seen in the wild after pdf-lib stamped a form.
    head, content = page_content("Misfiled xref recovered")
    write_objstm_pdf(
        "misfiled_xref.pdf",
        packed=[
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 6 0 R"
            b" /Resources << /Font << /F1 4 0 R >> /XObject << /Im1 7 0 R >> >> >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Producer (oxide-api fixtures) >>",
        ],
        loose=[
            ("1.1", b"1 1 obj\n<< /Type /Catalog /Pages 2 0 R /PageMode /UseNone >>", None),
            (6, b"6 0 obj\n" + head, content),
            (7, b"7 0 obj\n<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray"
                b" /BitsPerComponent 8 /Length 1 >>", b"\x80"),
        ],
        rows=lambda at: [(0, 0, 65535), (2, 8, 0), (1, at["1.1"], 1)]  # the stray row
        + [(2, 8, i) for i in range(1, 5)] + [(1, at[n], 0) for n in (6, 7, 8, 9)],
        trailer="/Root 1 1 R /Info 5 0 R",
    )

    # Recovering a misfiled object must not resurrect superseded ones. Page 3
    # was rewritten outside the ObjStm (its xref row is right); the Pages row
    # points at another object's header, so finding Pages means sweeping the
    # ObjStm, which still holds the old page 3.
    old_head, old_content = page_content("Superseded page text")
    new_head, new_content = page_content("Current page text")
    page = (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R"
            b" /Resources << /Font << /F1 4 0 R >> /XObject << /Im1 7 0 R >> >> >>")
    write_objstm_pdf(
        "misfiled_xref_superseded.pdf",
        packed=[
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            page % 5,
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ],
        loose=[
            (5, b"5 0 obj\n" + old_head, old_content),
            (7, b"7 0 obj\n<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray"
                b" /BitsPerComponent 8 /Length 1 >>", b"\x80"),
        ],
        update=[
            (6, b"6 0 obj\n" + new_head, new_content),
            ("3.new", b"3 0 obj\n" + page % 6, None),
        ],
        rows=lambda at: [(0, 0, 65535), (2, 8, 0), (1, at[6], 0), (1, at["3.new"], 0), (2, 8, 3)]
        + [(1, at[n], 0) for n in (5, 6, 7, 8, 9)],
        trailer="/Root 1 0 R",
    )


# --------------------------------------------------------------------------
# Images (sent straight to OCR)
# --------------------------------------------------------------------------

EN_LINES = ["Scanned Invoice 4471", "Total amount due: 1,250.00", "Payment terms: thirty days",
            "Thank you for your business"]
A4_200_TOP = (1654, 900)  # the top of an A4 page at 200 DPI: same text size, smaller files


def make_images():
    en = render_lines(EN_LINES, FONT_LATIN, 56, *A4_200_TOP)
    en.save(out("scan_en.png"), optimize=True)

    # A phone photo held sideways: the pixels are stored rotated 90° CCW and
    # EXIF Orientation 6 says "rotate 90° CW to display".
    exif = Image.Exif()
    exif[0x0112] = 6
    en.rotate(90, expand=True).convert("RGB").save(out("scan_en_exif.jpg"), quality=85, exif=exif)

    # A two-page bilevel fax-style TIFF (CCITT Group 4).
    p2 = render_lines(["Delivery note 5512", "Second page of the fax", "Marker tiff-beta"],
                      FONT_LATIN, 56, *A4_200_TOP)
    en.convert("1").save(out("scan_pages.tif"), save_all=True, append_images=[p2.convert("1")],
                         compression="group4", dpi=(200, 200))

    render_lines(["مرحبا بالعالم", "السلام عليكم ورحمة الله", "شكرا لكم"],
                 FONT_ARAB, 60, *A4_200_TOP, rtl=True).save(out("scan_ar.webp"), lossless=True)
    render_lines(["भारत का संविधान", "हम भारत के लोग", "न्याय स्वतंत्रता समता"],
                 FONT_DEVA, 60, *A4_200_TOP).save(out("scan_hi.jp2"))

    # Black text on a transparent background; transparent pixels stored black.
    line = render_lines(["Transparent label 8812"], FONT_LATIN, 56, 1000, 200, margin=50)
    rgba = Image.new("RGBA", line.size, (0, 0, 0, 0))
    rgba.putalpha(Image.eval(line, lambda v: 255 - v))
    rgba.save(out("transparent.png"), optimize=True)

    small = render_lines(["Format check 3141"], FONT_LATIN, 48, 900, 160, margin=40)
    small.save(out("line.gif"))
    small.save(out("line.bmp"))
    small.save(out("line.tga"))


# --------------------------------------------------------------------------
# OCR golden images (embedded in the binary for the startup smoke test)
# --------------------------------------------------------------------------

def make_golden():
    render_lines(["Invoice 2026 total due"], FONT_LATIN, 44, 760, 140, margin=40).save(f"{GOLDEN}/en.png", optimize=True)
    render_lines(["भारत का संविधान"], FONT_DEVA, 48, 560, 150, margin=40).save(f"{GOLDEN}/hi.png", optimize=True)
    render_lines(["مرحبا بالعالم"], FONT_ARAB, 48, 520, 150, rtl=True, margin=40).save(f"{GOLDEN}/ar.png", optimize=True)


def make_native_pdfs():
    make_native_pdf()
    lo_pdf_from_docx("native_hi.pdf", "हिंदी दस्तावेज़", ["भारत एक विशाल देश है।", "यह एक मूल पाठ वाला दस्तावेज़ है।"])
    lo_pdf_from_docx("native_ar.pdf", "وثيقة عربية", ["مرحبا بالعالم", "هذه وثيقة نصية أصلية"])
    make_arabic_forms_pdf()


GROUPS = {
    "office": make_office,
    "odf": make_odf,          # needs the office group's output
    "pdf": make_native_pdfs,
    "scans": make_scans,
    "overlay": make_scan_with_overlay,
    "ocr_layer": make_scans_with_ocr_layer,
    "hybrid": make_hybrid,
    "legacy": make_legacy_font_pdf,
    "encrypted": make_encrypted_pdfs,  # needs native.pdf
    "nested": make_nested_pdf,
    "misfiled": make_misfiled_xref_pdfs,
    "images": make_images,
    "golden": make_golden,
}

if __name__ == "__main__":
    for name in sys.argv[1:] or GROUPS:
        GROUPS[name]()
    for f in sorted(os.listdir(OUT)):
        print(f"{f:28s} {os.path.getsize(out(f)):>9,d} bytes")
